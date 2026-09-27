"""Integration tests for the Results Engine within the end-to-end PipelineOrchestrator.

Verifies:
QUALIFY -> ORDER -> SHOW -> query.limit -> RESULTS
"""

import io
import openpyxl
import pytest

from yogi.core.configuration.query import ScreeningQuery
from yogi.core.contracts.filter import FilterGroupSpec, FilterOperator, FilterRuleSpec
from yogi.core.contracts.ranking import RankingOrder
from yogi.core.contracts.results import ExportFormat, QueryResultPayload
from yogi.core.contracts.selection import RankCutoffSpec
from yogi.core.orchestration.pipeline import PipelineOrchestrator
from yogi.core.registry.hub import RegistryHub
from yogi.filter.engine import FilterEngine
from yogi.ranking.engine import RankingEngine
from yogi.results.assembler import ResultAssembler
from yogi.results.formatter import ExportFormatter
from yogi.selection.engine import SelectionEngine


@pytest.fixture
def sample_market_dataset():
    """Sample dataset of 5 stocks with varied metrics for end-to-end testing."""
    return {
        "ALPHA": {
            "symbol": "ALPHA",
            "name": "Alpha Ltd",
            "exchange": "NSE",
            "price": 100.0,
            "M_CUR": 25.0,
            "EMA_W": "upcut",
            "BANDST_M": "ABOVE",
            "S4": 15.0,
        },
        "BETA": {
            "symbol": "BETA",
            "name": "Beta Ltd",
            "exchange": "NSE",
            "price": 200.0,
            "M_CUR": 15.0,
            "EMA_W": "upcut",
            "BANDST_M": "IN",
            "S4": 10.0,
        },
        "GAMMA": {
            "symbol": "GAMMA",
            "name": "Gamma Ltd",
            "exchange": "NSE",
            "price": 300.0,
            "M_CUR": 30.0,
            "EMA_W": "downcut",
            "BANDST_M": "IN",
            "S4": 20.0,
        },
        "DELTA": {
            "symbol": "DELTA",
            "name": "Delta Ltd",
            "exchange": "NSE",
            "price": 400.0,
            "M_CUR": 5.0,
            "EMA_W": "upcut",
            "BANDST_M": "IN",
            "S4": 5.0,
        },
        "EPSILON": {
            "symbol": "EPSILON",
            "name": "Epsilon Ltd",
            "exchange": "NSE",
            "price": 500.0,
            "M_CUR": 18.0,
            "EMA_W": "upcut",
            "BANDST_M": "IN",
            "S4": 12.0,
        },
    }


def test_end_to_end_pipeline_with_result_assembler(sample_market_dataset):
    """Test 1: Full pipeline execution QUALIFY -> ORDER -> SHOW -> query.limit -> RESULTS using real ResultAssembler."""
    assembler = ResultAssembler()
    orchestrator = PipelineOrchestrator(
        registry=RegistryHub(),
        filter_rule=FilterEngine(),
        ranking_rule=RankingEngine(),
        selection_rule=SelectionEngine(),
        result_assembler=assembler,
    )

    # Screening query:
    # 1. QUALIFY: EMA_W == 'upcut' (Matches ALPHA, BETA, DELTA, EPSILON - 4 stocks)
    # 2. ORDER: Ranking by M_CUR DESC (ALPHA=25, EPSILON=18, BETA=15, DELTA=5)
    # 3. SHOW: Cutoff RANK_M_CUR <= 3 (ALPHA=1, EPSILON=2, BETA=3 survive - 3 stocks)
    # 4. limit: 2 (Top 2 survive: ALPHA, EPSILON)
    query = ScreeningQuery(
        filter_group=FilterGroupSpec(
            rules=[FilterRuleSpec(field="EMA_W", op=FilterOperator.EQ, value="upcut")]
        ),
        ranking_metric="M_CUR",
        ranking_order=RankingOrder.DESCENDING,
        selection_cutoffs=[RankCutoffSpec(rank_metric="RANK_M_CUR", max_rank=3)],
        limit=2,
    )

    payload = orchestrator.execute_screen(query=query, records=sample_market_dataset)

    # Verify QueryResultPayload structure
    assert isinstance(payload, QueryResultPayload)
    assert payload.universe_count == 5
    assert payload.qualified_count == 4
    assert payload.selected_count == 3
    assert len(payload.rows) == 2  # Limited to 2

    # Verify order is strictly preserved: ALPHA (rank 1), EPSILON (rank 2)
    assert payload.rows[0]["symbol"] == "ALPHA"
    assert payload.rows[0]["RANK_M_CUR"] == 1
    assert payload.rows[0]["M_CUR"] == 25.0

    assert payload.rows[1]["symbol"] == "EPSILON"
    assert payload.rows[1]["RANK_M_CUR"] == 2
    assert payload.rows[1]["M_CUR"] == 18.0

    # Verify column definitions include dynamic rank column
    col_keys = [c.key for c in payload.columns]
    assert "symbol" in col_keys
    assert "RANK_M_CUR" in col_keys
    assert "M_CUR" in col_keys


def test_pipeline_results_payload_export_to_xlsx_and_csv(sample_market_dataset):
    """Test 2: Pipeline query result payload can be directly formatted to valid XLSX and CSV bytes."""
    assembler = ResultAssembler()
    formatter = ExportFormatter()

    orchestrator = PipelineOrchestrator(
        registry=RegistryHub(),
        filter_rule=FilterEngine(),
        ranking_rule=RankingEngine(),
        selection_rule=SelectionEngine(),
        result_assembler=assembler,
    )

    query = ScreeningQuery(
        ranking_metric="M_CUR",
        ranking_order=RankingOrder.DESCENDING,
        limit=3,
    )

    payload = orchestrator.execute_screen(query=query, records=sample_market_dataset)
    assert len(payload.rows) == 3

    # Format to XLSX
    xlsx_bytes = formatter.format(payload, export_format=ExportFormat.XLSX)
    assert isinstance(xlsx_bytes, bytes)
    assert len(xlsx_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Screener"]
    assert ws.cell(row=1, column=1).value == "Symbol"
    assert ws.cell(row=2, column=1).value == "GAMMA"  # GAMMA has highest M_CUR=30

    # Format to CSV
    csv_bytes = formatter.format(payload, export_format=ExportFormat.CSV)
    assert isinstance(csv_bytes, bytes)
    assert b"GAMMA" in csv_bytes
