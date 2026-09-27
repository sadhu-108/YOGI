"""Tests for ExportFormatter — CSV, JSON, and XLSX In-Memory Serialization."""

import csv
import io
import json
import openpyxl
import pytest

from yogi.core.contracts.results import ColumnDefinition, ExportFormat, QueryResultPayload
from yogi.results.assembler import ResultAssembler
from yogi.results.formatter import ExportFormatter


@pytest.fixture
def sample_payload() -> QueryResultPayload:
    """Fixture providing a realistic assembled QueryResultPayload."""
    assembler = ResultAssembler()
    records = [
        {
            "symbol": "RELIANCE",
            "name": "Reliance Industries",
            "exchange": "NSE",
            "price": 2500.5,
            "RANK_M_CUR": 1,
            "M_CUR": 5.4,
            "BANDST_M": "IN",
            "sector": "Energy",
        },
        {
            "symbol": "TCS",
            "name": "Tata Consultancy",
            "exchange": "NSE",
            "price": 3500.0,
            "RANK_M_CUR": 2,
            "M_CUR": 4.1,
            "BANDST_M": "ABOVE",
            "sector": "IT",
        },
        {
            "symbol": "INFY",
            "name": "Infosys",
            "exchange": "NSE",
            "price": None,
            "RANK_M_CUR": 3,
            "M_CUR": None,
            "BANDST_M": None,
            "sector": "IT",
        },
    ]
    cols = ["symbol", "name", "price", "RANK_M_CUR", "M_CUR", "BANDST_M", "sector"]
    metadata = {"universe_count": 500, "qualified_count": 50, "selected_count": 3, "limit": 10}
    return assembler.assemble(records=records, columns=cols, metadata=metadata)


def test_formatter_csv_validity(sample_payload: QueryResultPayload):
    """Test 1: ExportFormatter formats valid RFC 4180 CSV bytes with empty string for None."""
    formatter = ExportFormatter()
    csv_bytes = formatter.format(payload=sample_payload, export_format=ExportFormat.CSV)

    assert isinstance(csv_bytes, bytes)
    csv_text = csv_bytes.decode("utf-8")

    # Read back with Python csv module
    reader = list(csv.reader(io.StringIO(csv_text)))
    assert len(reader) == 4  # 1 header + 3 rows

    # Verify header
    headers = reader[0]
    expected_labels = [c.label for c in sample_payload.columns]
    assert headers == expected_labels

    # Verify row 1
    assert reader[1][0] == "RELIANCE"
    assert reader[1][2] == "2500.5"
    assert reader[1][3] == "1"

    # Verify row 3 has empty strings for None values (NOT string 'None')
    assert reader[3][0] == "INFY"
    assert reader[3][2] == ""  # price was None
    assert reader[3][4] == ""  # M_CUR was None
    assert reader[3][5] == ""  # BANDST_M was None


def test_formatter_json_validity_and_absence_of_nan(sample_payload: QueryResultPayload):
    """Test 2: ExportFormatter formats valid JSON bytes with null for None and strictly zero NaN tokens."""
    formatter = ExportFormatter()
    json_bytes = formatter.format(payload=sample_payload, export_format=ExportFormat.JSON)

    assert isinstance(json_bytes, bytes)
    json_text = json_bytes.decode("utf-8")

    # Verify absence of NaN
    assert "NaN" not in json_text

    # Parse JSON
    parsed = json.loads(json_text)
    assert parsed["universe_count"] == 500
    assert parsed["qualified_count"] == 50
    assert parsed["selected_count"] == 3
    assert len(parsed["rows"]) == 3
    assert len(parsed["columns"]) == 7

    # Verify INFY row has null
    infy_row = parsed["rows"][2]
    assert infy_row["symbol"] == "INFY"
    assert infy_row["price"] is None
    assert infy_row["M_CUR"] is None


def test_formatter_xlsx_validity_and_styling(sample_payload: QueryResultPayload):
    """Test 3: ExportFormatter produces valid XLSX workbook bytes with dark header styling and freeze panes."""
    formatter = ExportFormatter()
    xlsx_bytes = formatter.format(payload=sample_payload, export_format=ExportFormat.XLSX)

    assert isinstance(xlsx_bytes, bytes)
    assert len(xlsx_bytes) > 0

    # Load workbook back from in-memory bytes
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    assert "Screener" in wb.sheetnames
    ws = wb["Screener"]

    # Verify freeze panes at B2
    assert ws.freeze_panes == "B2"

    # Verify auto-filter enabled
    assert ws.auto_filter.ref is not None

    # Verify header row styling (Row 1)
    cell_a1 = ws.cell(row=1, column=1)
    assert cell_a1.value == "Symbol"
    assert cell_a1.font.bold is True
    # Verify dark header fill color (1F2430) and gold font color (FFB055)
    assert cell_a1.fill.start_color.rgb == "001F2430" or cell_a1.fill.start_color.rgb == "1F2430"
    assert cell_a1.font.color.rgb == "00FFB055" or cell_a1.font.color.rgb == "FFB055"

    # Verify data rows (Row 2 to 4)
    cell_rel = ws.cell(row=2, column=1)
    assert cell_rel.value == "RELIANCE"

    cell_rank = ws.cell(row=2, column=4)  # RANK_M_CUR column
    assert cell_rank.value == 1
    assert cell_rank.number_format == "0"

    cell_price = ws.cell(row=2, column=3)  # price column
    assert cell_price.value == 2500.5
    assert cell_price.number_format == "0.00"

    # Verify None row has None in cell
    cell_infy_price = ws.cell(row=4, column=3)
    assert cell_infy_price.value is None


def test_formatter_empty_payload():
    """Test 4: ExportFormatter formats an empty payload cleanly across CSV, JSON, and XLSX without error."""
    formatter = ExportFormatter()
    empty_payload = QueryResultPayload(
        universe_count=0,
        qualified_count=0,
        selected_count=0,
        rows=[],
        columns=[
            ColumnDefinition(key="symbol", label="Symbol", category="IDENTIFIER", is_numeric=False),
            ColumnDefinition(key="price", label="Price", category="PRICE", is_numeric=True),
        ],
        metadata={},
    )

    # 1. Empty CSV (header only)
    csv_bytes = formatter.format(empty_payload, export_format=ExportFormat.CSV)
    csv_text = csv_bytes.decode("utf-8")
    lines = [l for l in csv_text.splitlines() if l]
    assert len(lines) == 1
    assert lines[0] == "Symbol,Price"

    # 2. Empty JSON
    json_bytes = formatter.format(empty_payload, export_format=ExportFormat.JSON)
    parsed = json.loads(json_bytes.decode("utf-8"))
    assert parsed["rows"] == []
    assert parsed["selected_count"] == 0

    # 3. Empty XLSX
    xlsx_bytes = formatter.format(empty_payload, export_format=ExportFormat.XLSX)
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb["Screener"]
    assert ws.cell(row=1, column=1).value == "Symbol"
    assert ws.cell(row=1, column=2).value == "Price"
    assert ws.cell(row=2, column=1).value is None


def test_formatter_unsupported_format_raises_value_error(sample_payload: QueryResultPayload):
    """Test 5: ExportFormatter raises ValueError on unsupported export formats."""
    formatter = ExportFormatter()

    with pytest.raises(ValueError, match="Unsupported export format"):
        formatter.format(payload=sample_payload, export_format="unsupported_format")  # type: ignore[arg-type]


def test_formatter_completely_empty_payload_csv():
    """Test 6: When both columns and rows are empty, CSV export returns a valid empty CSV representation without crashing or returning zero bytes."""
    formatter = ExportFormatter()
    empty_payload = QueryResultPayload(
        universe_count=0,
        qualified_count=0,
        selected_count=0,
        rows=[],
        columns=[],
        metadata={},
    )
    csv_bytes = formatter.format(empty_payload, export_format=ExportFormat.CSV)
    assert isinstance(csv_bytes, bytes)
    assert len(csv_bytes) > 0  # Must not silently return zero bytes
    # Valid CSV representation parseable by csv.reader without crash
    parsed = list(csv.reader(io.StringIO(csv_bytes.decode("utf-8"))))
    assert len(parsed) >= 1
