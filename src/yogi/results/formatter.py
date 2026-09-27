"""Export Formatter — Serializes QueryResultPayload into in-memory XLSX, CSV, and JSON bytes."""

import csv
import io
import json
import math
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from yogi.core.contracts.results import ExportFormat, IExportFormatter, QueryResultPayload


def _sanitize_for_json(obj: Any) -> Any:
    """Recursively sanitize data structures for JSON compliance (NaN -> None)."""
    if obj is None:
        return None
    if isinstance(obj, float) and math.isnan(obj):
        return None
    if isinstance(obj, dict):
        return {k: _sanitize_for_json(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [_sanitize_for_json(x) for x in obj]
    return obj


class ExportFormatter(IExportFormatter):
    """Serializes QueryResultPayload presentation objects into exportable binary or text bytes.

    Strictly produces in-memory bytes without modifying the payload, without mutating
    upstream data, and without writing to local files on disk.
    """

    def format(
        self,
        payload: QueryResultPayload,
        export_format: ExportFormat = ExportFormat.XLSX,
    ) -> bytes:
        """Format the payload into export bytes (XLSX, CSV, or JSON)."""
        fmt = export_format.value if isinstance(export_format, ExportFormat) else str(export_format).lower()

        if fmt == ExportFormat.CSV.value:
            return self._format_csv(payload)
        elif fmt == ExportFormat.JSON.value:
            return self._format_json(payload)
        elif fmt == ExportFormat.XLSX.value:
            return self._format_xlsx(payload)
        else:
            raise ValueError(f"Unsupported export format: {export_format!r}. Supported formats: XLSX, CSV, JSON.")

    def _format_csv(self, payload: QueryResultPayload) -> bytes:
        """Serialize payload rows into standard RFC 4180 CSV bytes."""
        columns = payload.columns
        rows = payload.rows

        # If columns metadata is available, use column keys and labels
        if columns:
            keys = [col.key for col in columns]
            headers = [col.label for col in columns]
        elif rows:
            keys = list(rows[0].keys())
            headers = keys
        else:
            keys = []
            headers = []

        buf = io.StringIO()
        writer = csv.writer(buf, lineterminator="\r\n")

        # Write header if available, otherwise write a valid empty record line
        if headers:
            writer.writerow(headers)
        else:
            writer.writerow([])

        # Write data rows
        for row in rows:
            writer.writerow([row.get(k, "") if row.get(k) is not None else "" for k in keys])

        return buf.getvalue().encode("utf-8")

    def _format_json(self, payload: QueryResultPayload) -> bytes:
        """Serialize payload into formatted, UTF-8 JSON bytes with NaN sanitized to null."""
        data = {
            "universe_count": payload.universe_count,
            "qualified_count": payload.qualified_count,
            "selected_count": payload.selected_count,
            "columns": [
                {
                    "key": col.key,
                    "label": col.label,
                    "category": col.category,
                    "is_numeric": col.is_numeric,
                    "format_string": col.format_string,
                }
                for col in payload.columns
            ],
            "rows": payload.rows,
            "metadata": payload.metadata,
        }

        sanitized = _sanitize_for_json(data)
        json_str = json.dumps(sanitized, indent=2, allow_nan=False, default=str)
        return json_str.encode("utf-8")

    def _format_xlsx(self, payload: QueryResultPayload) -> bytes:
        """Serialize payload into a styled, professional Excel XLSX workbook in memory."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Screener"

        columns = payload.columns
        rows = payload.rows

        # Determine effective keys and labels
        if columns:
            keys = [col.key for col in columns]
            labels = [col.label for col in columns]
            formats = [col.format_string for col in columns]
        elif rows:
            keys = list(rows[0].keys())
            labels = keys
            formats = [None for _ in keys]
        else:
            # Completely empty payload
            keys = []
            labels = []
            formats = []

        # Dark header styles matching verified prototype aesthetic
        hdr_fill = PatternFill(start_color="1F2430", end_color="1F2430", fill_type="solid")
        hdr_font = Font(name="Calibri", size=10, bold=True, color="FFB055")
        thin_side = Side(style="thin", color="3A4150")
        hdr_border = Border(top=thin_side, bottom=thin_side, left=thin_side, right=thin_side)
        body_font = Font(name="Calibri", size=10)

        # 1. Write Header Row
        for c_idx, label in enumerate(labels, start=1):
            cell = ws.cell(row=1, column=c_idx, value=label)
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = hdr_border

            # Set proportional column width
            col_letter = get_column_letter(c_idx)
            ws.column_dimensions[col_letter].width = max(9, min(24, len(str(label)) + 4))

        # 2. Write Data Rows
        for r_idx, row in enumerate(rows, start=2):
            for c_idx, key in enumerate(keys, start=1):
                val = row.get(key)
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = body_font

                if isinstance(val, (int, float)) and not isinstance(val, bool):
                    fmt_str = formats[c_idx - 1]
                    if fmt_str:
                        cell.number_format = fmt_str
                    elif key.startswith("RANK_"):
                        cell.number_format = "0"
                    else:
                        cell.number_format = "0.00"

        # 3. Apply Sheet Usability Settings
        if labels:
            ws.freeze_panes = "B2"  # Freeze header and first column (symbol)

        if rows and labels:
            ws.auto_filter.ref = ws.dimensions  # Enable Excel dropdown filter arrows

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()
